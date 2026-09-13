"""Lectura aislada de CSV/XLSX para staging; no persiste ni confirma."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import load_workbook

from talentia.shared.application.errores import EntradaInvalidaError


def leer_filas(nombre: str, contenido: bytes, maximo: int = 5000) -> list[dict[str, object]]:
    extension = Path(nombre).suffix.casefold()
    if extension == ".csv":
        try:
            texto = contenido.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise EntradaInvalidaError("El CSV debe usar UTF-8") from exc
        filas = [dict(fila) for fila in csv.DictReader(io.StringIO(texto))]
    elif extension == ".xlsx":
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hoja = libro.active
        valores = hoja.iter_rows(values_only=True)
        encabezados = [str(valor or "").strip() for valor in next(valores, ())]
        filas = [dict(zip(encabezados, fila, strict=False)) for fila in valores]
        libro.close()
    else:
        raise EntradaInvalidaError("Solo se admiten archivos CSV o XLSX")
    if not filas or len(filas) > maximo:
        raise EntradaInvalidaError("El lote debe contener entre 1 y 5000 filas")
    return filas
