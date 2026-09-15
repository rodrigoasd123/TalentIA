"""Generador de 10 CVs en formato DOCX profesionales para probar la carga masiva y el sistema de alertas de TalentIA."""

from pathlib import Path
import zipfile
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT_DIR = Path(__file__).resolve().parent.parent
DEST_DIR = ROOT_DIR / "descargas_talento" / "05_Lote_Pruebas_10_CVs"
DEST_DIR.mkdir(parents=True, exist_ok=True)


CASOS_CV = [
    {
        "archivo": "01_CV_Gonzalo_Benavides_Apto_Limpio.docx",
        "nombre": "Gonzalo Javier Benavides Rivas",
        "dni": "72109845",
        "correo": "gonzalo.benavides@gmail.com",
        "telefono": "+51 991122334",
        "ubicacion": "Miraflores, Lima",
        "fuente": "LinkedIn",
        "convocatoria": "DEV-BACK-01 - Desarrollador Backend Senior (Python / Cloud)",
        "salario": "S/ 8,200.00",
        "disponibilidad": "Inmediata",
        "skills": "Python, FastAPI, Docker, PostgreSQL, AWS, CI/CD, Microservicios",
        "caso_uso": "Candidato 100% Apto / Limpio (Sin antecedentes en bases TCS)",
        "alerta_esperada": "Sin alertas (Apto / Registrado)",
        "resumen": "Ingeniero de Software con más de 6 años de experiencia en desarrollo backend con Python y arquitecturas basadas en nube AWS.",
    },
    {
        "archivo": "02_CV_Mariana_Quiroz_Apto_Limpia.docx",
        "nombre": "Mariana Alejandra Quiroz Tello",
        "dni": "73450912",
        "correo": "mariana.quiroz@gmail.com",
        "telefono": "+51 982345678",
        "ubicacion": "San Isidro, Lima",
        "fuente": "Computrabajo",
        "convocatoria": "QA-AUTO-02 - Ingeniero QA Automation (Web & API)",
        "salario": "S/ 6,800.00",
        "disponibilidad": "15 dias",
        "skills": "Selenium, Playwright, Cypress, Python, PyTest, Postman, Jenkins",
        "caso_uso": "Candidata 100% Apta / Limpia para QA Automation",
        "alerta_esperada": "Sin alertas (Apto / Registrado)",
        "resumen": "Especialista en Aseguramiento de Calidad con enfoque en automatización de pruebas end-to-end y pruebas de integración de microservicios.",
    },
    {
        "archivo": "03_CV_Carlos_Mendoza_ExTCS_Elegible.docx",
        "nombre": "Carlos Alberto Mendoza Ramos",
        "dni": "45871234",
        "correo": "carlos.mendoza.exp@gmail.com",
        "telefono": "+51 987451234",
        "ubicacion": "Santiago de Surco, Lima",
        "fuente": "Referido Interno",
        "convocatoria": "DEV-BACK-01 - Desarrollador Backend Senior (Python / Cloud)",
        "salario": "S/ 8,500.00",
        "disponibilidad": "Inmediata",
        "skills": "Python, Django, FastAPI, Kubernetes, AWS, Terraform, Clean Architecture",
        "caso_uso": "Excolaborador TCS con salida voluntaria y calificación sobresaliente (Elegible para Reingreso)",
        "alerta_esperada": "Alerta Informativa: Coincidencia con excolaborador TCS elegible para reingreso",
        "resumen": "Desarrollador Senior que laboró anteriormente en proyectos de Banca Digital en TCS. Renuncia voluntaria formal por oportunidad en el exterior.",
    },
    {
        "archivo": "04_CV_Lucia_Vargas_ExTCS_Elegible.docx",
        "nombre": "Lucia Mariana Vargas Silva",
        "dni": "46982345",
        "correo": "lucia.vargas.tech@gmail.com",
        "telefono": "+51 983456789",
        "ubicacion": "San Borja, Lima",
        "fuente": "LinkedIn",
        "convocatoria": "DEV-BACK-01 - Desarrollador Backend Senior (Python / Cloud)",
        "salario": "S/ 9,000.00",
        "disponibilidad": "Inmediata",
        "skills": "Java 17, Spring Boot, Microservicios, Kafka, Oracle, Docker, AWS",
        "caso_uso": "Excolaboradora TCS con salida por cierre de proyecto (Calificación A - Excelente)",
        "alerta_esperada": "Alerta Informativa: Coincidencia con excolaborador TCS elegible para reingreso",
        "resumen": "Líder técnico con destacada trayectoria en TCS liderando squads de automatización para clientes de telecomunicaciones.",
    },
    {
        "archivo": "05_CV_Martin_Rojas_ExTCS_NoElegible_y_Vetado.docx",
        "nombre": "Martin Gonzalo Rojas Paredes",
        "dni": "49345678",
        "correo": "martin.rojas.dev@gmail.com",
        "telefono": "+51 984567890",
        "ubicacion": "San Miguel, Lima",
        "fuente": "Bumeran",
        "convocatoria": "DEV-BACK-01 - Desarrollador Backend Senior (Python / Cloud)",
        "salario": "S/ 7,500.00",
        "disponibilidad": "Inmediata",
        "skills": "Java, Spring Boot, PostgreSQL, Microservicios, Docker, Git",
        "caso_uso": "Excolaborador TCS NO Elegible (Falta grave por abandono de puesto) y Vetado Permanente",
        "alerta_esperada": "Doble Alerta Alta: Excolaborador TCS no elegible + Restricción de veto vigente",
        "resumen": "Desarrollador Java con experiencia en aplicaciones transaccionales bancarias y microservicios empresariales.",
    },
    {
        "archivo": "06_CV_Jorge_Navarro_ExTCS_NoElegible_Carencia.docx",
        "nombre": "Jorge Luis Navarro Salgado",
        "dni": "51567890",
        "correo": "jorge.navarro.ops@gmail.com",
        "telefono": "+51 985678901",
        "ubicacion": "Lince, Lima",
        "fuente": "Indeed",
        "convocatoria": "DEV-BACK-01 - Desarrollador Backend Senior (Python / Cloud)",
        "salario": "S/ 7,200.00",
        "disponibilidad": "Inmediata",
        "skills": "DevOps, Kubernetes, Docker, Linux, Bash, Python, AWS",
        "caso_uso": "Excolaborador TCS No Elegible (No superó período de prueba técnico) con período de enfriamiento",
        "alerta_esperada": "Alerta Alta: Excolaborador TCS no elegible (Revisión obligatoria)",
        "resumen": "Ingeniero de infraestructura y automatización con foco en despliegues sobre Kubernetes y contenedores Docker.",
    },
    {
        "archivo": "07_CV_Camila_Caceres_Vetada_Confidencialidad.docx",
        "nombre": "Camila Andrea Caceres Torres",
        "dni": "54890123",
        "correo": "camila.caceres.qa@gmail.com",
        "telefono": "+51 986789012",
        "ubicacion": "Surquillo, Lima",
        "fuente": "LinkedIn",
        "convocatoria": "QA-AUTO-02 - Ingeniero QA Automation (Web & API)",
        "salario": "S/ 7,000.00",
        "disponibilidad": "Inmediata",
        "skills": "Selenium, Playwright, Python, Postman, SQL, Git, Jira",
        "caso_uso": "Postulante Vetada de TCS (Veto Permanente por violación grave de confidencialidad bancaria)",
        "alerta_esperada": "Alerta Crítica / Alta: Coincidencia con una restricción vigente (Veto Permanente)",
        "resumen": "Ingeniera de Pruebas con experiencia en banca y retail, especializada en pruebas automatizadas y funcionales.",
    },
    {
        "archivo": "08_CV_Braulio_Paredes_Vetado_TituloAdulterado.docx",
        "nombre": "Braulio Marcelo Paredes Quispe",
        "dni": "71829304",
        "correo": "braulio.paredes.data@gmail.com",
        "telefono": "+51 987890123",
        "ubicacion": "Chorrillos, Lima",
        "fuente": "Portal Web",
        "convocatoria": "DATA-ENG-03 - Ingeniero de Datos (ETL & Big Data)",
        "salario": "S/ 8,000.00",
        "disponibilidad": "Inmediata",
        "skills": "Python, Spark, SQL, Airflow, Snowflake, AWS Glue, ETL",
        "caso_uso": "Postulante Vetado Permanente (Título y certificados adulterados detectados en verificación BGC)",
        "alerta_esperada": "Alerta Crítica / Alta: Coincidencia con una restricción vigente (Veto Permanente)",
        "resumen": "Profesional en ingeniería enfocado en pipelines de datos distribuidos y modelado analítico empresarial.",
    },
    {
        "archivo": "09_CV_Javier_Villalobos_Vetado_ConflictoInteres.docx",
        "nombre": "Javier Eduardo Villalobos Ruiz",
        "dni": "34125678",
        "correo": "javier.villalobos.arch@gmail.com",
        "telefono": "+51 988901234",
        "ubicacion": "Magdalena del Mar, Lima",
        "fuente": "Headhunter",
        "convocatoria": "DEV-BACK-01 - Desarrollador Backend Senior (Python / Cloud)",
        "salario": "S/ 9,500.00",
        "disponibilidad": "Inmediata",
        "skills": "Microservicios, Cloud Architecture, Python, Go, Kubernetes, Kafka",
        "caso_uso": "Postulante con Inhabilitación Ético-Legal Activa (Conflicto de interés y competencia desleal comprobada)",
        "alerta_esperada": "Alerta Alta: Coincidencia con restricción activa de cumplimiento",
        "resumen": "Arquitecto de Soluciones Cloud con más de 10 años diseñando plataformas escalables y microservicios de alto tráfico.",
    },
    {
        "archivo": "10_CV_Duplicado_Gonzalo_Benavides.docx",
        "nombre": "Gonzalo Javier Benavides Rivas",
        "dni": "72109845",
        "correo": "gonzalo.benavides@gmail.com",
        "telefono": "+51 991122334",
        "ubicacion": "Miraflores, Lima",
        "fuente": "Bumeran",
        "convocatoria": "DEV-BACK-01 - Desarrollador Backend Senior (Python / Cloud)",
        "salario": "S/ 8,400.00",
        "disponibilidad": "Inmediata",
        "skills": "Python, FastAPI, Docker, PostgreSQL, AWS, Kubernetes, Terraform",
        "caso_uso": "Candidato Duplicado (Mismo DNI que el CV 01) para validar que no cree duplicados en BD",
        "alerta_esperada": "Estado: REUTILIZADO (No crea doble ficha, conserva histórico y agrega documento)",
        "resumen": "Segunda postulación del mismo candidato con pretensión actualizada. Debe asociarse a la ficha existente.",
    },
]


def crear_documento_cv(datos: dict) -> Path:
    doc = docx.Document()

    # Formato de página
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # 1. Nombre completo (Línea 1 limpia)
    p_nombre = doc.add_paragraph()
    r_nombre = p_nombre.add_run(datos["nombre"])
    r_nombre.font.size = Pt(20)
    r_nombre.font.bold = True
    r_nombre.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    # 2. Datos de contacto estructurados
    p_contacto = doc.add_paragraph()
    p_contacto.add_run(f"DNI: {datos['dni']} | Correo: {datos['correo']} | Celular: {datos['telefono']}\n")
    p_contacto.add_run(f"Ubicacion: {datos['ubicacion']} | Canal de atraccion: {datos['fuente']}\n")
    p_contacto.add_run(f"Convocatoria: {datos['convocatoria']} | Expectativa salarial: {datos['salario']}\n")
    p_contacto.add_run(f"Disponibilidad: {datos['disponibilidad']} | Especialidad: {datos['skills']}")

    # Cuadro informativo del caso de prueba
    p_caso = doc.add_paragraph()
    r_caso_label = p_caso.add_run(f"CASO DE PRUEBA: {datos['caso_uso']}\n")
    r_caso_label.font.bold = True
    r_caso_label.font.size = Pt(10)
    r_caso_label.font.color.rgb = RGBColor(0x85, 0x64, 0x04)
    r_alerta = p_caso.add_run(f"RESPUESTA ESPERADA DEL SISTEMA: {datos['alerta_esperada']}")
    r_alerta.font.italic = True
    r_alerta.font.size = Pt(9.5)

    doc.add_paragraph("─" * 45)

    # 3. Resumen Ejecutivo
    h1 = doc.add_heading("1. Resumen Ejecutivo", level=2)
    h1.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    doc.add_paragraph(datos["resumen"])

    # 4. Habilidades Técnicas
    h2 = doc.add_heading("2. Competencias y Habilidades Técnicas", level=2)
    h2.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    p_skills = doc.add_paragraph()
    p_skills.add_run(f"Skills Principales: {datos['skills']}\n")
    p_skills.add_run("Metodologías: Scrum, Agile, CI/CD, Git Flow, Clean Code, TDD")

    # 5. Experiencia Laboral
    h3 = doc.add_heading("3. Experiencia Laboral Relevante", level=2)
    h3.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    p_exp = doc.add_paragraph()
    p_exp.add_run("Especialista Técnico Senior | Empresa Líder de Tecnología (2021 - 2024)\n").bold = True
    p_exp.add_run("• Diseño e implementación de servicios de misión crítica con alta concurrencia.\n")
    p_exp.add_run("• Coordinación con equipos multidisciplinarios y cumplimiento de estándares de calidad.\n")
    p_exp.add_run("• Optimización continua de código y pipelines de integración automatizada.")

    # 6. Educación
    h4 = doc.add_heading("4. Educación y Certificaciones", level=2)
    h4.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    p_edu = doc.add_paragraph()
    p_edu.add_run("• Titulado en Ingeniería de Sistemas / Computación e Informática\n")
    p_edu.add_run("• Certificación Internacional en Cloud Computing / DevOps / QA")

    ruta_salida = DEST_DIR / datos["archivo"]
    doc.save(ruta_salida)
    return ruta_salida


def generar_todos():
    print(f"Generando 10 CVs en: {DEST_DIR}")
    archivos_generados = []
    for caso in CASOS_CV:
        ruta = crear_documento_cv(caso)
        archivos_generados.append(ruta)
        print(f"  [OK] Generado: {caso['archivo']}")

    # Generar archivo ZIP en descargas_talento y en raíz del proyecto
    zip_ruta_1 = DEST_DIR.parent / "05_Lote_Pruebas_10_CVs.zip"
    zip_ruta_2 = ROOT_DIR / "TalentIA_Lote_10_CVs_Prueba_Alertas.zip"

    for zpath in (zip_ruta_1, zip_ruta_2):
        with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
            for fpath in archivos_generados:
                z.write(fpath, arcname=fpath.name)
        print(f"  [ZIP] Paquete ZIP creado: {zpath.name}")

    # Generar README explicativo
    readme_path = DEST_DIR / "GUIA_DE_PRUEBAS_10_CVS.md"
    with open(readme_path, "w", encoding="utf-8") as f:
        f.write("# Guía de Casos de Prueba - Lote de 10 CVs para TalentIA\n\n")
        f.write("Este lote ha sido diseñado específicamente para validar de forma exhaustiva los motores de extracción automática y alertas de control de TalentIA.\n\n")
        f.write("| N° | Archivo | Candidato | DNI | Caso de Uso | Comportamiento / Alerta Esperada |\n")
        f.write("|---|---|---|---|---|---|\n")
        for i, c in enumerate(CASOS_CV, 1):
            f.write(f"| {i} | `{c['archivo']}` | {c['nombre']} | `{c['dni']}` | {c['caso_uso']} | **{c['alerta_esperada']}** |\n")
        f.write("\n\n### Instrucciones para probar en la Web:\n")
        f.write("1. Ingresa a `http://127.0.0.1:8000/candidatos/importar-cvs`\n")
        f.write("2. Selecciona varios o todos los archivos de este directorio.\n")
        f.write("3. Presiona **Importar currículums**.\n")
        f.write("4. Observa cómo el sistema autocompleta la base general y levanta las alertas amarillas y rojas con los antecedentes.\n")
        f.write("5. Sube el archivo `10_CV_Duplicado_Gonzalo_Benavides.docx` para comprobar que el sistema detecta la identidad exacta y lo marca como **reutilizado** sin crear duplicados.\n")

    print(f"  [DOC] Guia de pruebas generada: {readme_path.name}")


if __name__ == "__main__":
    generar_todos()
