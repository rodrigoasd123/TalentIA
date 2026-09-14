import os
import sys
from pathlib import Path
import docx

ROOT_DIR = Path('.').resolve()
sys.path.insert(0, str(ROOT_DIR))
sys.path.insert(0, str(ROOT_DIR / 'src'))

from scripts.seed_convocatorias_candidatos import CONVOCATORIAS, CANDIDATOS_DATOS

DEST_DIR = ROOT_DIR / 'descargas_talento'
CONV_DIR = DEST_DIR / '01_Convocatorias'
CVS_DIR = DEST_DIR / '02_Curriculums_25_Postulantes'

CONV_DIR.mkdir(parents=True, exist_ok=True)
CVS_DIR.mkdir(parents=True, exist_ok=True)

# 1. Generar convocatorias
for i, conv in enumerate(CONVOCATORIAS, 1):
    cod = conv['codigo']
    titulo = conv['titulo']
    ctc = float(conv['ctc'])
    reqs = conv['requisitos']
    
    clean_title = titulo.replace(' ', '_').replace('/', '_').replace('(', '').replace(')', '')
    filename_base = f"{i:02d}_{cod}_{clean_title}"
    
    # Markdown
    lines = [
        f"# CONVOCATORIA LABORAL: {titulo}",
        "",
        f"**Código de Proceso:** `{cod}`  ",
        "**Empresa Cliente:** Tata Consultancy Services (TCS) - Perú  ",
        f"**Presupuesto Máximo (CTC Mensual):** S/ {ctc:,.2f}  ",
        "**Modalidad de Trabajo:** Híbrido / Remoto  ",
        "**Ubicación:** Lima / Arequipa / Trujillo, Perú  ",
        "",
        "---",
        "",
        "## 1. Descripción del Puesto",
        f"Estamos en la búsqueda de un/a **{titulo}** para integrarse a proyectos estratégicos de transformación digital en clientes del sector corporativo.",
        "",
        "## 2. Requisitos del Perfil",
        "",
        "### Requisitos Obligatorios (Excluyentes):"
    ]
    for r in reqs:
        if r['obligatorio']:
            lines.append(f"- **[{r['codigo']}]**: {r['descripcion']} (Ponderación: {r['peso']})")
    
    lines.append("")
    lines.append("### Requisitos Deseables (No excluyentes):")
    for r in reqs:
        if not r['obligatorio']:
            lines.append(f"- **[{r['codigo']}]**: {r['descripcion']} (Ponderación: {r['peso']})")
            
    lines.extend([
        "",
        "## 3. Beneficios Corporativos",
        "- Planilla completa directa con todos los beneficios de ley.",
        "- Seguro médico EPS cubierto al 100% para el colaborador.",
        "- Bonos por desempeño y certificaciones oficiales en Cloud / AI.",
        "- Acceso a la plataforma global de aprendizaje de TCS."
    ])
    
    with open(CONV_DIR / f"{filename_base}.md", "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
        
    # Word .docx
    doc = docx.Document()
    doc.add_heading(titulo, level=1)
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Código: {cod} | Cliente: Tata Consultancy Services (TCS)\n").bold = True
    p_meta.add_run(f"Compensación Total (CTC): S/ {ctc:,.2f} mensuales\nModalidad: Híbrido / Remoto (Perú)")
    
    doc.add_heading("1. Descripción del Puesto", level=2)
    doc.add_paragraph(f"Estamos en la búsqueda de un/a {titulo} para integrarse a proyectos estratégicos de escala empresarial.")
    
    doc.add_heading("2. Requisitos Obligatorios", level=2)
    for r in reqs:
        if r['obligatorio']:
            doc.add_paragraph(f"• [{r['codigo']}] {r['descripcion']} (Ponderación: {r['peso']})", style="List Bullet")
            
    doc.add_heading("3. Requisitos Deseables", level=2)
    for r in reqs:
        if not r['obligatorio']:
            doc.add_paragraph(f"• [{r['codigo']}] {r['descripcion']} (Ponderación: {r['peso']})", style="List Bullet")
            
    doc.add_heading("4. Beneficios Corporativos", level=2)
    doc.add_paragraph("• Planilla directa con todos los beneficios de ley.")
    doc.add_paragraph("• Seguro médico EPS cubierto al 100%.")
    doc.add_paragraph("• Capacitación constante y plan de certificaciones oficiales.")
    
    doc.save(CONV_DIR / f"{filename_base}.docx")

print(f"Convocatorias generadas exitosamente: {len(CONVOCATORIAS)}")

# 2. Generar 25 Curriculums
UNIVERSIDADES = [
    'Pontificia Universidad Católica del Perú (PUCP)',
    'Universidad Nacional de Ingeniería (UNI)',
    'Universidad de Lima',
    'Universidad Peruana de Ciencias Aplicadas (UPC)',
    'Universidad Nacional Mayor de San Marcos (UNMSM)',
    'Universidad Nacional de San Agustín (UNSA - Arequipa)'
]

EMPRESAS_TI = [
    ('Fintech Andina S.A.C.', 'Tech Lead / Desarrollador Especialista'),
    ('Banco Continental Digital', 'Ingeniero de Software Senior'),
    ('Software Factory Latam', 'Consultor Técnico'),
    ('Telefónica Tech Perú', 'Especialista en Automatización e Infraestructura'),
    ('Retail Corp Solutions', 'Ingeniero de Datos y Procesos')
]

for i, c_data in enumerate(CANDIDATOS_DATOS, 1):
    (nombres, apellidos, doc_raw, correo, tel_raw, f_nac, ubicacion, fuente, conocimiento, expectativa, estado, conv_idx) = c_data
    conv = CONVOCATORIAS[conv_idx]
    
    nombre_completo = f"{nombres} {apellidos}"
    clean_name = nombre_completo.replace(" ", "_")
    filename_base = f"{i:02d}_CV_{clean_name}"
    uni = UNIVERSIDADES[(i * 3) % len(UNIVERSIDADES)]
    emp1, rol1 = EMPRESAS_TI[(i * 2) % len(EMPRESAS_TI)]
    emp2, rol2 = EMPRESAS_TI[(i * 2 + 1) % len(EMPRESAS_TI)]
    sueldo = float(expectativa)
    
    # Markdown
    cv_md_lines = [
        f"# CURRICULUM VITAE - {nombre_completo.upper()}",
        "",
        f"**DNI / Documento:** {doc_raw}  ",
        f"**Correo Electrónico:** [{correo}](mailto:{correo})  ",
        f"**Teléfono / Celular:** +51 {tel_raw}  ",
        f"**Ubicación:** {ubicacion}  ",
        f"**Canal de Atracción:** {fuente}  ",
        f"**Convocatoria de Postulación:** `{conv['codigo']}` - {conv['titulo']}  ",
        f"**Expectativa Salarial:** S/ {sueldo:,.2f}  ",
        f"**Estado en Proceso:** `{estado.upper()}`  ",
        "",
        "---",
        "",
        "## 1. Resumen Ejecutivo",
        f"Profesional en Ingeniería de Sistemas con sólida trayectoria en {conocimiento}. Apasionado por la excelencia técnica, metodologías ágiles y soluciones que agregan valor directo al negocio.",
        "",
        "## 2. Competencias Técnicas",
        f"- **Especialidad:** {conocimiento}",
        "- **Bases de Datos:** PostgreSQL, MySQL, SQLite, MongoDB, Redis",
        "- **Herramientas & Cloud:** Git, Docker, Kubernetes, CI/CD, AWS / GCP",
        "- **Buenas Prácticas:** Clean Architecture, SOLID, TDD, Microservicios",
        "",
        "## 3. Experiencia Laboral",
        "",
        f"### **{rol1}** | *{emp1}*",
        "*2022 - Presente | Lima, Perú*",
        "- Diseño y despliegue de componentes clave para servicios transaccionales de alta demanda.",
        "- Implementación de pruebas automatizadas y pipelines de integración continua con GitHub Actions.",
        "- Optimización de consultas de base de datos y rendimiento de servicios RESTful.",
        "",
        f"### **{rol2}** | *{emp2}*",
        "*2019 - 2022 | Perú*",
        "- Desarrollo y mantenimiento de arquitecturas de software modulares.",
        "- Integración de APIs internas y externas para sincronización de datos en tiempo real.",
        "- Trabajo colaborativo con equipos multidisciplinarios de producto y QA.",
        "",
        "## 4. Educación y Formación",
        f"- **Bachiller / Titulado en Ingeniería de Sistemas y Computación**  ",
        f"  *{uni}* (2013 - 2018)",
        "",
        "## 5. Idiomas y Certificaciones",
        "- **Español:** Nativo",
        "- **Inglés:** Nivel Profesional / Intermedio Avanzado (B2)",
        "- **Certificaciones:** Certificación Oficial de Especialidad / Scrum Master"
    ]
    
    with open(CVS_DIR / f"{filename_base}.md", "w", encoding="utf-8") as f:
        f.write("\n".join(cv_md_lines))
        
    # Word .docx
    doc = docx.Document()
    doc.add_heading(nombre_completo, level=1)
    p_info = doc.add_paragraph()
    p_info.add_run(f"DNI: {doc_raw} | Celular: +51 {tel_raw}\n").bold = True
    p_info.add_run(f"Correo: {correo} | Ubicación: {ubicacion}\n")
    p_info.add_run(f"Postulación: {conv['codigo']} - {conv['titulo']} | Expectativa: S/ {sueldo:,.2f}")
    
    doc.add_heading("1. Resumen Ejecutivo", level=2)
    doc.add_paragraph(
        f"Ingeniero/a con amplia experiencia en el sector tecnológico y foco en {conocimiento}. "
        "Enfocado/a en resultados, arquitectura limpia y soluciones de alto impacto para proyectos de misión crítica."
    )
    
    doc.add_heading("2. Habilidades Técnicas", level=2)
    doc.add_paragraph(f"• Especialidad: {conocimiento}", style="List Bullet")
    doc.add_paragraph("• Metodologías y Herramientas: Git, Docker, CI/CD, Scrum, Clean Architecture", style="List Bullet")
    
    doc.add_heading("3. Experiencia Laboral", level=2)
    p_exp1 = doc.add_paragraph()
    p_exp1.add_run(f"{rol1} – {emp1} (2022 - Presente)\n").bold = True
    p_exp1.add_run("• Responsable del desarrollo, mantenimiento y evolución de arquitecturas de software empresariales.\n")
    p_exp1.add_run("• Automatización de pipelines y soporte a requerimientos de alto volumen.")
    
    p_exp2 = doc.add_paragraph()
    p_exp2.add_run(f"{rol2} – {emp2} (2019 - 2022)\n").bold = True
    p_exp2.add_run("• Implementación de integraciones RESTful y modernización hacia microservicios.")
    
    doc.add_heading("4. Educación y Certificaciones", level=2)
    doc.add_paragraph(f"• {uni} – Grado en Ingeniería")
    doc.add_paragraph("• Idiomas: Español (Nativo), Inglés (Intermedio - Avanzado)")
    
    doc.save(CVS_DIR / f"{filename_base}.docx")

print(f"Curriculums generados exitosamente: {len(CANDIDATOS_DATOS)}")