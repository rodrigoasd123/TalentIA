"""Lectura defensiva de CSV/XLSX sin ejecutar contenido de las celdas."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from app.core.exceptions import ValidationError

CANONICAL_FIELDS = frozenset({
    "full_name", "email", "phone", "national_id", "location", "source",
    "job_code", "historical_status", "applied_at", "legal_basis_status",
    "linkedin_url",
})

ALIASES: dict[str, tuple[str, ...]] = {
    "full_name": ("nombre", "nombres", "nombre completo", "candidato", "full name"),
    "email": ("email", "correo", "correo electronico", "e-mail"),
    "phone": ("telefono", "celular", "movil", "phone"),
    "national_id": ("documento", "dni", "cedula", "document number", "id"),
    "location": ("ubicacion", "ciudad", "location"),
    "source": ("fuente", "origen", "source", "proveedor"),
    "job_code": ("vacante", "codigo vacante", "requisicion", "rgs", "job code"),
    "historical_status": ("estado", "etapa", "status"),
    "applied_at": ("fecha postulacion", "fecha", "applied at"),
    "legal_basis_status": ("base legal", "consentimiento", "privacy status"),
    "linkedin_url": ("linkedin", "linkedin url", "perfil linkedin"),
}


def normalize_header(value: str) -> str:
    import unicodedata

    text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text.strip().lower().replace("_", " "))


@dataclass(slots=True)
class TabularData:
    headers: list[str]
    rows: list[dict[str, str]]
    sheet_name: str = ""
    sheet_names: list[str] | None = None


class FileSafetyScanner(Protocol):
    """Punto de integración para un escáner antimalware aprobado."""

    def scan(self, content: bytes, filename: str) -> None: ...


class TabularReader:
    def __init__(
        self,
        *,
        max_bytes: int,
        max_rows: int,
        safety_scanner: FileSafetyScanner | None = None,
    ) -> None:
        self.max_bytes = max_bytes
        self.max_rows = max_rows
        self.safety_scanner = safety_scanner

    def read(self, content: bytes, filename: str, *, sheet_name: str = "") -> TabularData:
        if not content:
            raise ValidationError("El archivo está vacío")
        if len(content) > self.max_bytes:
            raise ValidationError("El archivo excede el límite configurado")
        if self.safety_scanner is not None:
            self.safety_scanner.scan(content, filename)
        suffix = Path(filename).suffix.lower()
        if suffix == ".csv":
            return self._read_csv(content)
        if suffix == ".xlsx":
            return self._read_xlsx(content, sheet_name=sheet_name)
        raise ValidationError("Solo se permiten archivos CSV o XLSX")

    def _read_csv(self, content: bytes) -> TabularData:
        if b"\x00" in content[:4096]:
            raise ValidationError("El archivo no parece ser un CSV de texto")
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except UnicodeDecodeError as exc:
                raise ValidationError("No se pudo decodificar el CSV") from exc
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text, newline=""), dialect)
        try:
            headers = [str(value).strip() for value in next(reader)]
        except StopIteration as exc:
            raise ValidationError("El CSV no contiene encabezados") from exc
        self._validate_headers(headers)
        rows = []
        for number, values in enumerate(reader, start=2):
            if number - 1 > self.max_rows:
                raise ValidationError("El archivo excede el máximo de filas configurado")
            padded = values + [""] * (len(headers) - len(values))
            rows.append(
                {header: str(padded[index]).strip() for index, header in enumerate(headers)}
            )
        return TabularData(headers=headers, rows=rows)

    def _read_xlsx(self, content: bytes, *, sheet_name: str) -> TabularData:
        if not content.startswith(b"PK"):
            raise ValidationError("El contenido no corresponde a un XLSX válido")
        self._validate_xlsx_archive(content)
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True, keep_links=False
            )
        except Exception as exc:
            raise ValidationError("No se pudo abrir el XLSX; puede estar cifrado o dañado") from exc
        names = list(workbook.sheetnames)
        selected = sheet_name or (names[0] if names else "")
        if selected not in names:
            workbook.close()
            raise ValidationError("La hoja solicitada no existe")
        sheet = workbook[selected]
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value or "").strip() for value in next(iterator)]
        except StopIteration as exc:
            workbook.close()
            raise ValidationError("La hoja no contiene encabezados") from exc
        self._validate_headers(headers)
        rows = []
        for number, values in enumerate(iterator, start=2):
            if number - 1 > self.max_rows:
                workbook.close()
                raise ValidationError("El archivo excede el máximo de filas configurado")
            rows.append(
                {
                    header: str(
                        values[index]
                        if index < len(values) and values[index] is not None
                        else ""
                    ).strip()
                    for index, header in enumerate(headers)
                }
            )
        workbook.close()
        return TabularData(headers=headers, rows=rows, sheet_name=selected, sheet_names=names)

    def _validate_xlsx_archive(self, content: bytes) -> None:
        """Rechaza contenedores activos o con expansión desproporcionada."""
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > 5_000:
                    raise ValidationError("El XLSX contiene demasiados elementos internos")
                expanded_limit = max(self.max_bytes * 20, 100 * 1024 * 1024)
                expanded_size = sum(member.file_size for member in members)
                if expanded_size > expanded_limit:
                    raise ValidationError("El XLSX excede el límite de expansión segura")
                lowered_names = {member.filename.lower() for member in members}
                unsafe_parts = (
                    "vbaproject.bin",
                    "externallinks/",
                    "embeddings/",
                    "oleobjects/",
                )
                if any(part in name for name in lowered_names for part in unsafe_parts):
                    raise ValidationError(
                        "El XLSX contiene macros, enlaces externos u objetos no admitidos"
                    )
                if any(".." in Path(name).parts for name in lowered_names):
                    raise ValidationError("El XLSX contiene rutas internas no seguras")
        except zipfile.BadZipFile as exc:
            raise ValidationError("El contenido no corresponde a un XLSX válido") from exc

    @staticmethod
    def _validate_headers(headers: list[str]) -> None:
        if not headers or any(not header for header in headers):
            raise ValidationError("Todos los encabezados deben tener nombre")
        normalized = [normalize_header(header) for header in headers]
        if len(set(normalized)) != len(normalized):
            raise ValidationError("El archivo contiene encabezados duplicados")

    @staticmethod
    def suggest_mapping(headers: list[str]) -> dict[str, str]:
        suggestions: dict[str, str] = {}
        normalized = {header: normalize_header(header) for header in headers}
        for canonical, aliases in ALIASES.items():
            for header, value in normalized.items():
                if value in aliases or value == canonical.replace("_", " "):
                    suggestions[canonical] = header
                    break
        return suggestions


def neutralize_spreadsheet_formula(value: str) -> str:
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


__all__ = [
    "CANONICAL_FIELDS", "FileSafetyScanner", "TabularData", "TabularReader",
    "neutralize_spreadsheet_formula", "normalize_header",
]
