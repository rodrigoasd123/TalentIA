"""SPEC-005: importación histórica gobernada, idempotente y autorizada."""

from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.application.services.audit_service import Actor
from app.application.services.email_service import EmailService
from app.application.unit_of_work import UnitOfWork
from app.application.use_cases.historical_import import HistoricalImportUseCase
from app.core.config import reset_settings_cache
from app.core.exceptions import ConsentMissingOrExpired, ValidationError
from app.domain.entities import Candidate, EmailTemplate, Job
from app.domain.enums import EmailTemplateKind, JobStatus, Role
from app.domain.imports import ImportStatus, RowClassification
from app.infrastructure.database.models import Base
from app.infrastructure.database.session import get_session_factory, reset_engine
from app.infrastructure.imports.tabular_reader import TabularReader, neutralize_spreadsheet_formula
from app.infrastructure.repositories.sqlalchemy_repos import SqlApplicationRepository
from app.infrastructure.security.tokens import TokenService

pytestmark = pytest.mark.integration


@pytest.fixture
def import_uow(monkeypatch, tmp_path):
    monkeypatch.setenv("TALENTIA_IMPORT_STORAGE_PATH", str(tmp_path / "imports"))
    monkeypatch.setenv("TALENTIA_IMPORT_MAX_FILE_MB", "20")
    monkeypatch.setenv("TALENTIA_IMPORT_MAX_ROWS", "10000")
    reset_settings_cache()
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield UnitOfWork(session)
    reset_settings_cache()


def _csv(*rows: str) -> bytes:
    return ("Nombre,Correo,Telefono,Documento,Vacante\n" + "\n".join(rows)).encode()


def _open_job(uow: UnitOfWork, code: str = "VAC-IMP") -> Job:
    job = Job(code=code, title="Vacante ficticia", status=JobStatus.OPEN)
    return uow.jobs.add(job)


def test_importa_candidato_restringido_y_es_idempotente(import_uow: UnitOfWork) -> None:
    _open_job(import_uow)
    actor = Actor.user("hr-manager")
    service = HistoricalImportUseCase(import_uow)
    content = _csv("Persona Histórica,historica@example.test,,,VAC-IMP")

    batch, reused = service.upload(
        content=content, filename="historico.csv", source="archivo", actor=actor
    )
    assert reused is False
    repeated, reused = service.upload(
        content=content, filename="otro-nombre.csv", source="archivo", actor=actor
    )
    assert reused is True
    assert repeated.id == batch.id

    batch = service.validate(
        batch_id=batch.id,
        mapping={
            "full_name": "Nombre",
            "email": "Correo",
            "phone": "Telefono",
            "national_id": "Documento",
            "job_code": "Vacante",
        },
        actor=actor,
        expected_version=batch.version,
    )
    assert batch.status is ImportStatus.READY_FOR_REVIEW
    confirmed = service.confirm(
        batch_id=batch.id,
        confirmation_key="confirmacion-especifica-001",
        actor=actor,
        expected_version=batch.version,
    )
    assert confirmed.status is ImportStatus.IMPORTED
    assert confirmed.summary["created_candidates"] == 1
    candidate = import_uow.candidates.get_by_email("historica@example.test")
    assert candidate is not None
    assert candidate.processing_status == "restricted_review"
    assert candidate.legal_basis_status == "unknown"
    assert candidate.can_be_processed is False
    applications = import_uow.applications.list_for_candidate(candidate.id)
    assert len(applications) == 1

    template = import_uow.templates.add(
        EmailTemplate(
            code="historico_restringido",
            kind=EmailTemplateKind.RECEIPT_CONFIRMATION,
            subject_template="Prueba {job_title}",
            body_template="Mensaje de prueba",
            approved=True,
        )
    )
    with pytest.raises(ConsentMissingOrExpired, match="restringido"):
        EmailService(import_uow).prepare(
            application_id=applications[0].id,
            template_code=template.code,
            actor=actor,
            use_ai=False,
        )

    repeated_confirmation = service.confirm(
        batch_id=batch.id,
        confirmation_key="confirmacion-especifica-001",
        actor=actor,
    )
    assert repeated_confirmation.id == batch.id
    assert len(import_uow.applications.list_for_candidate(candidate.id)) == 1


def test_nombre_solo_permanece_en_staging_para_revision(import_uow: UnitOfWork) -> None:
    _open_job(import_uow)
    service = HistoricalImportUseCase(import_uow)
    actor = Actor.user("recruiter")
    batch, _ = service.upload(
        content=_csv("Nombre Ambiguo,,,,VAC-IMP"),
        filename="ambiguo.csv",
        source="archivo",
        actor=actor,
    )
    batch = service.validate(
        batch_id=batch.id,
        mapping={"full_name": "Nombre", "job_code": "Vacante"},
        actor=actor,
    )
    assert batch.status is ImportStatus.REJECTED
    rows = import_uow.imports.list_rows(batch.id)
    assert rows[0].classification is RowClassification.MANUAL_REVIEW_REQUIRED
    assert import_uow.candidates.list() == []


def test_varios_lotes_sin_confirmar_no_comparten_clave_vacia(
    import_uow: UnitOfWork,
) -> None:
    service = HistoricalImportUseCase(import_uow)
    actor = Actor.user("recruiter")
    first, _ = service.upload(
        content=_csv("Persona Uno,uno@example.test,,,VAC-UNO"),
        filename="uno.csv",
        source="archivo",
        actor=actor,
    )
    second, _ = service.upload(
        content=_csv("Persona Dos,dos@example.test,,,VAC-DOS"),
        filename="dos.csv",
        source="archivo",
        actor=actor,
    )
    assert first.id != second.id
    assert first.confirmation_key is None
    assert second.confirmation_key is None


def test_telefono_normalizado_es_identificador_fuerte(import_uow: UnitOfWork) -> None:
    _open_job(import_uow)
    existing = import_uow.candidates.add(
        Candidate(full_name="Persona Teléfono", phone="+51 999-123-456")
    )
    service = HistoricalImportUseCase(import_uow)
    batch, _ = service.upload(
        content=_csv("Persona Telefono,,999123456,,VAC-IMP"),
        filename="telefono.csv",
        source="archivo",
        actor=Actor.user("recruiter"),
    )
    service.validate(
        batch_id=batch.id,
        mapping={
            "full_name": "Nombre",
            "phone": "Telefono",
            "job_code": "Vacante",
        },
        actor=Actor.user("recruiter"),
    )
    row = import_uow.imports.list_rows(batch.id)[0]
    assert row.classification is RowClassification.EXACT_DUPLICATE
    assert row.candidate_id == existing.id


def test_lector_xlsx_no_ejecuta_formulas_y_rechaza_encabezados_duplicados() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Nombre", "Vacante"])
    sheet.append(["=HYPERLINK(\"https://invalid.test\")", "VAC-001"])
    stream = io.BytesIO()
    workbook.save(stream)
    parsed = TabularReader(max_bytes=1_000_000, max_rows=10).read(
        stream.getvalue(), "entrada.xlsx"
    )
    assert parsed.rows[0]["Nombre"] == ""
    assert neutralize_spreadsheet_formula("=1+1") == "'=1+1"

    with pytest.raises(ValidationError, match="duplicados"):
        TabularReader(max_bytes=1000, max_rows=10).read(
            b"Nombre,nombre\nA,B", "duplicado.csv"
        )


def test_lector_xlsx_rechaza_macros_y_objetos_activos() -> None:
    workbook = Workbook()
    workbook.active.append(["Nombre", "Vacante"])
    stream = io.BytesIO()
    workbook.save(stream)
    hostile = io.BytesIO(stream.getvalue())
    with zipfile.ZipFile(hostile, mode="a") as archive:
        archive.writestr("xl/vbaProject.bin", b"contenido-ficticio")

    with pytest.raises(ValidationError, match="macros"):
        TabularReader(max_bytes=1_000_000, max_rows=10).read(
            hostile.getvalue(), "activo.xlsx"
        )


def test_lector_invoca_punto_de_escaneo_antimalware() -> None:
    class RejectingScanner:
        def scan(self, content: bytes, filename: str) -> None:
            assert content.startswith(b"Nombre")
            assert filename == "muestra.csv"
            raise ValidationError("El escáner rechazó el archivo")

    reader = TabularReader(
        max_bytes=1_000,
        max_rows=10,
        safety_scanner=RejectingScanner(),
    )
    with pytest.raises(ValidationError, match="escáner rechazó"):
        reader.read(b"Nombre,Vacante\nPersona,VAC-1", "muestra.csv")


def test_selecciona_otra_hoja_xlsx_antes_del_mapeo(import_uow: UnitOfWork) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Primera"
    first.append(["Nombre", "Vacante"])
    first.append(["Persona Primera", "VAC-UNO"])
    second = workbook.create_sheet("Segunda")
    second.append(["Candidato", "Requisicion"])
    second.append(["Persona Segunda", "VAC-DOS"])
    stream = io.BytesIO()
    workbook.save(stream)
    service = HistoricalImportUseCase(import_uow)
    actor = Actor.user("recruiter")
    batch, _ = service.upload(
        content=stream.getvalue(),
        filename="hojas.xlsx",
        source="archivo",
        actor=actor,
    )

    selected = service.select_sheet(
        batch_id=batch.id,
        sheet_name="Segunda",
        expected_version=batch.version,
        actor=actor,
    )

    assert selected.sheet_name == "Segunda"
    assert selected.summary["headers"] == ["Candidato", "Requisicion"]
    assert selected.suggested_mapping == {
        "full_name": "Candidato",
        "job_code": "Requisicion",
    }
    assert import_uow.imports.list_rows(batch.id)[0].raw_data["Candidato"] == "Persona Segunda"


@pytest.fixture
def api_client(monkeypatch, tmp_path):
    monkeypatch.setenv("TALENTIA_DATABASE_URL", f"sqlite:///{tmp_path / 'imports.db'}")
    monkeypatch.setenv("TALENTIA_IMPORT_STORAGE_PATH", str(tmp_path / "storage"))
    monkeypatch.setenv("TALENTIA_ENVIRONMENT", "development")
    reset_settings_cache()
    reset_engine()
    from app.api.main import app

    with TestClient(app) as client:
        yield client
    reset_engine()
    reset_settings_cache()


def test_confirmacion_exige_permiso_hr(api_client: TestClient) -> None:
    job = api_client.post(
        "/api/v1/jobs",
        json={"code": "VAC-API", "title": "Vacante API", "criteria_approved": True},
    ).json()
    assert job["status"] == "open"
    uploaded = api_client.post(
        "/api/v1/imports/historical",
        data={"source": "test"},
        files={
            "file": (
                "historico.csv",
                _csv("Persona API,api@example.test,,,VAC-API"),
                "text/csv",
            )
        },
    ).json()
    validated = api_client.post(
        f"/api/v1/imports/{uploaded['id']}/validate",
        json={
            "mapping": {
                "full_name": "Nombre",
                "email": "Correo",
                "job_code": "Vacante",
            },
            "expected_version": uploaded["version"],
        },
    ).json()
    denied = api_client.post(
        f"/api/v1/imports/{uploaded['id']}/confirm",
        json={"idempotency_key": "api-confirm-001", "expected_version": validated["version"]},
    )
    assert denied.status_code == 403

    auditor_token = TokenService().issue_pair(
        user_id="auditor-test", email="auditor@example.test", role=Role.AUDITOR
    ).access_token
    minimized = api_client.get(
        f"/api/v1/imports/{uploaded['id']}/rows",
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    assert minimized.status_code == 200
    assert minimized.json()["rows"][0]["data"]["email"] == "ap***"

    token = TokenService().issue_pair(
        user_id="manager-test", email="manager@example.test", role=Role.HIRING_MANAGER
    ).access_token
    allowed = api_client.post(
        f"/api/v1/imports/{uploaded['id']}/confirm",
        headers={"Authorization": f"Bearer {token}"},
        json={"idempotency_key": "api-confirm-001", "expected_version": validated["version"]},
    )
    assert allowed.status_code == 200, allowed.text
    assert allowed.json()["status"] == "imported"


def test_cancelacion_y_reporte_de_errores_son_gobernados(api_client: TestClient) -> None:
    uploaded = api_client.post(
        "/api/v1/imports/historical",
        data={"source": "test"},
        files={"file": ("errores.csv", _csv("Persona Error,error@, ,,NO-EXISTE"), "text/csv")},
    ).json()
    validated = api_client.post(
        f"/api/v1/imports/{uploaded['id']}/validate",
        json={
            "mapping": {
                "full_name": "Nombre",
                "email": "Correo",
                "job_code": "Vacante",
            }
        },
    )
    assert validated.status_code == 200
    report = api_client.get(f"/api/v1/imports/{uploaded['id']}/errors.csv")
    assert report.status_code == 200
    assert "error@" not in report.text
    assert "job_code" in report.text

    denied = api_client.post(
        f"/api/v1/imports/{uploaded['id']}/cancel",
        json={"reason": "Corrección del archivo de prueba"},
    )
    assert denied.status_code == 403
    token = TokenService().issue_pair(
        user_id="manager-test", email="manager@example.test", role=Role.HIRING_MANAGER
    ).access_token
    cancelled = api_client.post(
        f"/api/v1/imports/{uploaded['id']}/cancel",
        headers={"Authorization": f"Bearer {token}"},
        json={"reason": "Corrección del archivo de prueba"},
    )
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    with get_session_factory()() as session:
        uow = UnitOfWork(session)
        events = uow.audit.list(
            resource_id=uploaded["id"], action="import.cancelled"
        )
        assert len(events) == 1
        assert events[0].metadata["reason"] == "Corrección del archivo de prueba"
        assert uow.audit.verify_chain() == (True, None)


def test_fallo_durante_confirmacion_revierte_candidato(
    api_client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    api_client.post(
        "/api/v1/jobs",
        json={"code": "VAC-ROLLBACK", "title": "Vacante rollback", "criteria_approved": True},
    )
    uploaded = api_client.post(
        "/api/v1/imports/historical",
        data={"source": "test"},
        files={
            "file": (
                "rollback.csv",
                _csv("Persona Rollback,rollback@example.test,,,VAC-ROLLBACK"),
                "text/csv",
            )
        },
    ).json()
    validated = api_client.post(
        f"/api/v1/imports/{uploaded['id']}/validate",
        json={
            "mapping": {
                "full_name": "Nombre",
                "email": "Correo",
                "job_code": "Vacante",
            }
        },
    ).json()
    token = TokenService().issue_pair(
        user_id="manager-test", email="manager@example.test", role=Role.HIRING_MANAGER
    ).access_token

    def fail_application(*_args, **_kwargs):
        raise RuntimeError("fallo ficticio posterior a crear Candidate")

    monkeypatch.setattr(SqlApplicationRepository, "add", fail_application)
    with pytest.raises(RuntimeError, match="fallo ficticio"):
        api_client.post(
            f"/api/v1/imports/{uploaded['id']}/confirm",
            headers={"Authorization": f"Bearer {token}"},
            json={
                "idempotency_key": "rollback-confirm-001",
                "expected_version": validated["version"],
            },
        )

    with get_session_factory()() as session:
        uow = UnitOfWork(session)
        assert uow.candidates.get_by_email("rollback@example.test") is None
        batch = uow.imports.get_batch(uploaded["id"])
        assert batch is not None
        assert batch.status is ImportStatus.READY_FOR_REVIEW
