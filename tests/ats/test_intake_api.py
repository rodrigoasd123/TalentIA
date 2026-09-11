from __future__ import annotations

import io

import docx
import pytest
from fastapi.testclient import TestClient

from app.core.config import reset_settings_cache
from app.infrastructure.database.session import reset_engine
from app.application.unit_of_work import UnitOfWork
from app.infrastructure.database.session import session_scope


@pytest.fixture
def client(monkeypatch, tmp_path) -> TestClient:
    monkeypatch.setenv("TALENTIA_DATABASE_URL", f"sqlite:///{tmp_path / 'intake.db'}")
    monkeypatch.setenv("TALENTIA_ENVIRONMENT", "development")
    reset_settings_cache()
    reset_engine()
    from app.api.main import app

    with TestClient(app) as test_client:
        yield test_client
    reset_engine()
    reset_settings_cache()


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("CV ficticio de prueba", level=1)
    document.add_paragraph(
        "Analista con cinco años de experiencia en Python, SQL, automatización, "
        "integración de datos y elaboración de reportes para equipos de negocio."
    )
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_alta_completa_crea_vacante_candidato_cv_y_candidatura(client: TestClient) -> None:
    job = client.post(
        "/api/v1/jobs",
        json={
            "code": "LAB-001",
            "title": "Analista de datos",
            "description": "Vacante ficticia",
            "mandatory_skills": ["python", "sql"],
            "criteria_approved": True,
        },
    )
    assert job.status_code == 201
    assert job.json()["status"] == "open"

    response = client.post(
        "/api/v1/intake",
        data={
            "full_name": "Persona Ficticia",
            "email": "persona@example.test",
            "job_id": job.json()["id"],
            "consent_granted": "true",
            "source": "manual",
        },
        files={
            "resume": (
                "cv-ficticio.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["candidate_id"]
    assert payload["resume_id"]
    assert payload["application_id"]

    applications = client.get("/api/v1/applications")
    assert applications.status_code == 200
    assert any(item["id"] == payload["application_id"] for item in applications.json())


def test_ingesta_rechaza_falta_de_consentimiento(client: TestClient) -> None:
    job = client.post(
        "/api/v1/jobs",
        json={"code": "LAB-002", "title": "Vacante ficticia", "criteria_approved": True},
    ).json()
    response = client.post(
        "/api/v1/intake",
        data={
            "full_name": "Persona Sin Consentimiento",
            "email": "sin-consentimiento@example.test",
            "job_id": job["id"],
            "consent_granted": "false",
        },
        files={"resume": ("cv.docx", _docx_bytes(), "application/octet-stream")},
    )
    assert response.status_code == 422, response.text
    assert "consentimiento" in response.text.lower()


def test_editar_criterios_reabre_solo_con_aprobacion_humana(client: TestClient) -> None:
    job = client.post(
        "/api/v1/jobs",
        json={"code": "LAB-003", "title": "Vacante editable", "criteria_approved": True},
    ).json()
    edited = client.patch(
        f"/api/v1/jobs/{job['id']}", json={"mandatory_skills": ["python"]}
    )
    assert edited.status_code == 200
    assert edited.json()["status"] == "draft"
    approved = client.patch(
        f"/api/v1/jobs/{job['id']}", json={"criteria_approved": True}
    )
    assert approved.status_code == 200
    assert approved.json()["status"] == "open"


def test_sourcing_solo_devuelve_una_consulta_para_ejecucion_manual(client: TestClient) -> None:
    job = client.post(
        "/api/v1/jobs",
        json={"code": "LAB-004", "title": "Data Engineer", "mandatory_skills": ["python", "sql"], "criteria_approved": True},
    ).json()
    response = client.get(f"/api/v1/jobs/{job['id']}/sourcing-query")
    assert response.status_code == 200
    assert response.json()["execution"] == "manual"
    assert response.json()["opens_linkedin"] is False
    assert '"python"' in response.json()["boolean_query"]


def _create_application(client: TestClient, code: str = "LAB-022") -> dict:
    job = client.post(
        "/api/v1/jobs",
        json={"code": code, "title": "Vacante de integridad", "criteria_approved": True},
    ).json()
    response = client.post(
        "/api/v1/intake",
        data={
            "full_name": f"Persona {code}",
            "email": f"{code.lower()}@example.test",
            "job_id": job["id"],
            "consent_granted": "true",
        },
        files={"resume": ("cv.docx", _docx_bytes(), "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_lista_identifica_y_endpoint_asocia_cv_faltante(client: TestClient) -> None:
    created = _create_application(client)
    with session_scope() as database_session:
        uow = UnitOfWork(database_session)
        application = uow.applications.get(created["application_id"])
        assert application is not None
        application.resume_id = None
        uow.applications.update(application)

    before = client.get("/api/v1/applications").json()
    row = next(item for item in before if item["id"] == created["application_id"])
    assert row["has_resume"] is False
    assert row["resume_id"] is None

    attached = client.post(
        f"/api/v1/applications/{created['application_id']}/resume",
        files={"resume": ("cv-recuperado.docx", _docx_bytes(), "application/octet-stream")},
    )
    assert attached.status_code == 200, attached.text
    assert attached.json()["resume_id"] == created["resume_id"]

    after = client.get("/api/v1/applications").json()
    row = next(item for item in after if item["id"] == created["application_id"])
    assert row["has_resume"] is True


def test_solicitud_revision_manual_es_idempotente(client: TestClient) -> None:
    created = _create_application(client, "LAB-023")
    url = f"/api/v1/applications/{created['application_id']}/reviews"
    payload = {
        "reason": "criterion_unverified",
        "note": "Validar nivel de inglés por llamada.",
    }
    first = client.post(url, json=payload)
    second = client.post(url, json=payload)

    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    assert first.json()["created"] is True
    assert second.json()["created"] is False
    assert first.json()["id"] == second.json()["id"]

    rows = client.get("/api/v1/applications").json()
    row = next(item for item in rows if item["id"] == created["application_id"])
    assert row["has_open_review"] is True

    statistics = client.get("/api/v1/reviews/statistics").json()
    assert statistics["pending"] == 1
    assert "applications_in_review_state" in statistics
